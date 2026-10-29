"""
Dashboard de Folha Salarial (CLT + PJ).

Só ROLE_ADMIN acessa (administradores no Portal) — é dado financeiro sensível.
Não é um sistema de lançamento/edição: o usuário só importa a planilha
mensal (.xlsx) e o sistema calcula os indicadores. Reimportar a mesma
competência substitui os dados antigos dela (não duplica).

Planilha esperada (mesmo layout usado pelo RH hoje), 2 abas relevantes:
  "Detalhado Relatorio" — uma linha por colaborador CLT, com Regional e
      a coluna "Custo" (OPERACIONAL/ADMINISTRATIVO → mapeada para MOD/MOI).
  "PJ"                  — uma linha por prestador (Fornecedor, Salario, Descrição).
A aba "FINANCEIRO PAGAMENTO" é um resumo das mesmas pessoas e não é lida
(evita duplicar/divergir dos números de "Detalhado Relatorio").

Rotas:
  POST   /api/folha/importar                      — sobe o .xlsx de uma competência
  GET    /api/folha/importacoes                    — lista competências já importadas
  GET    /api/folha/resumo?importacao_id=          — indicadores da competência (default: mais recente)
  GET    /api/folha/evolucao-headcount?funcao=     — qtd de colaboradores por competência (mensal)
  GET    /api/folha/auditoria?importacao_id=       — variações de valor vs. a competência anterior (CLT+PJ)
  DELETE /api/folha/importacoes/<id>               — apaga uma importação (reimportar manda de novo)
"""
import io
from datetime import date, datetime

import openpyxl
from flask import Blueprint, request, jsonify
from sqlalchemy import func as sa_func, and_ as sa_and, or_ as sa_or

import audit
import models
from database import get_db
from security import require_role_admin

bp = Blueprint("folha", __name__, url_prefix="/api/folha")

ABA_DETALHADO = "Detalhado Relatorio"
ABA_PJ        = "PJ"

# cabeçalho (coluna → campo) da aba "Detalhado Relatorio", nessa ordem
COLUNAS_DETALHADO = [
    "nome", "funcao", "data_admissao", "salario_base", "periculosidade",
    "gratificacao", "ferias", "salario_familia", "cesta_basica", "plano_saude",
    "desconto_coparticipacao", "desconto_consignado", "adiantamento",
    "desconto_falta", "desconto_odr", "desconto_antecipacao_ferias",
    "inss", "irpf", "fgts", "ajuda_custo", "total_liquido",
    "perspectiva_salarial", "custo_raw", "regional", "contrato",
]

# Regra: tudo que é OPERACIONAL é MOD (mão de obra direta); qualquer outra
# coisa — ADMINISTRATIVO, em branco, ou qualquer outro valor — é MOI
# (mão de obra indireta). Não existe mais "Não informado" nessa classificação.
def _classificacao(custo_raw):
    return "MOD" if (custo_raw or "").upper() == "OPERACIONAL" else "MOI"

MESES_PT = [
    "", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
    "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro",
]


def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(".", "").replace(",", "."))
    except ValueError:
        return 0.0


def _txt(v):
    if v is None:
        return None
    s = str(v).replace("\xa0", " ").strip()
    return s or None


def _normaliza_nome(s):
    return (s or "").strip().upper()


def _parse_detalhado(ws):
    """Lê a aba 'Detalhado Relatorio' até a linha de total (primeira linha sem
    nome na coluna A depois de começarem os dados), devolve lista de dicts."""
    linhas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nome = _txt(row[0]) if len(row) > 0 else None
        if not nome:
            # primeira linha em branco = fim da tabela (linha de total/rodapé abaixo é ignorada)
            if linhas:
                break
            continue
        valores = list(row) + [None] * (len(COLUNAS_DETALHADO) - len(row))
        d = dict(zip(COLUNAS_DETALHADO, valores))

        data_adm = d["data_admissao"]
        if isinstance(data_adm, datetime):
            data_adm = data_adm.date()
        elif not isinstance(data_adm, date):
            data_adm = None

        classificacao = _classificacao(_txt(d["custo_raw"]))

        linhas.append({
            "nome": nome,
            "funcao": _txt(d["funcao"]),
            "data_admissao": data_adm,
            "regional": _txt(d["regional"]),
            "classificacao": classificacao,
            "salario_base": _num(d["salario_base"]),
            "periculosidade": _num(d["periculosidade"]),
            "gratificacao": _num(d["gratificacao"]),
            "ferias": _num(d["ferias"]),
            "salario_familia": _num(d["salario_familia"]),
            "cesta_basica": _num(d["cesta_basica"]),
            "plano_saude": _num(d["plano_saude"]),
            "desconto_coparticipacao": _num(d["desconto_coparticipacao"]),
            "desconto_consignado": _num(d["desconto_consignado"]),
            "adiantamento": _num(d["adiantamento"]),
            "desconto_falta": _num(d["desconto_falta"]),
            "desconto_odr": _num(d["desconto_odr"]),
            "desconto_antecipacao_ferias": _num(d["desconto_antecipacao_ferias"]),
            "inss": _num(d["inss"]),
            "irpf": _num(d["irpf"]),
            "fgts": _num(d["fgts"]),
            "ajuda_custo": _num(d["ajuda_custo"]),
            "total_liquido": _num(d["total_liquido"]),
            "perspectiva_salarial": _num(d["perspectiva_salarial"]),
            "contrato": _txt(d["contrato"]),
        })
    return linhas


def _parse_pj(ws):
    linhas = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        fornecedor = _txt(row[0]) if len(row) > 0 else None
        if not fornecedor:
            continue
        valor = _num(row[1]) if len(row) > 1 else 0.0
        descricao = _txt(row[2]) if len(row) > 2 else None
        setor = _txt(row[3]) if len(row) > 3 else None   # 4ª coluna, quando a planilha tiver (opcional)
        linhas.append({"fornecedor": fornecedor, "valor": valor, "descricao": descricao, "setor": setor})
    return linhas


@bp.route("/importar", methods=["POST"])
@require_role_admin
def importar():
    arquivo = request.files.get("arquivo")
    if not arquivo or not arquivo.filename:
        return jsonify({"message": "Envie o arquivo da planilha (.xlsx)"}), 400
    if not arquivo.filename.lower().endswith((".xlsx", ".xlsm")):
        return jsonify({"message": "Formato não suportado — envie um .xlsx"}), 400

    try:
        mes = int(request.form.get("mes", ""))
        ano = int(request.form.get("ano", ""))
        if not (1 <= mes <= 12) or not (2000 <= ano <= 2100):
            raise ValueError
    except ValueError:
        return jsonify({"message": "Informe mês (1-12) e ano válidos"}), 400

    try:
        conteudo = arquivo.read()
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True)
    except Exception:
        return jsonify({"message": "Não consegui abrir o arquivo — confira se é um .xlsx válido"}), 400

    if ABA_DETALHADO not in wb.sheetnames:
        return jsonify({"message": f"A planilha precisa ter a aba \"{ABA_DETALHADO}\""}), 400

    colaboradores = _parse_detalhado(wb[ABA_DETALHADO])
    pjs = _parse_pj(wb[ABA_PJ]) if ABA_PJ in wb.sheetnames else []

    if not colaboradores:
        return jsonify({"message": f"Não encontrei colaboradores na aba \"{ABA_DETALHADO}\""}), 400

    total_clt = round(sum(c["total_liquido"] for c in colaboradores), 2)
    total_pj  = round(sum(p["valor"] for p in pjs), 2)

    db = get_db()
    try:
        existente = (
            db.query(models.FolhaImportacao)
            .filter_by(competencia_mes=mes, competencia_ano=ano)
            .first()
        )
        if existente:
            db.delete(existente)
            db.flush()

        importacao = models.FolhaImportacao(
            competencia_mes=mes,
            competencia_ano=ano,
            arquivo_nome=arquivo.filename,
            importado_por=request.username,
            total_clt=total_clt,
            total_pj=total_pj,
            total_geral=round(total_clt + total_pj, 2),
            qtd_clt=len(colaboradores),
            qtd_pj=len(pjs),
        )
        db.add(importacao)
        db.flush()

        for c in colaboradores:
            db.add(models.FolhaColaborador(importacao_id=importacao.id, **c))
        for p in pjs:
            db.add(models.FolhaPJ(importacao_id=importacao.id, **p))

        db.commit()

        audit.log(
            request.username, "IMPORT_FOLHA", entity="folha_importacoes", entity_id=importacao.id,
            detail=f"{MESES_PT[mes]}/{ano} — {len(colaboradores)} CLT, {len(pjs)} PJ, "
                   f"total R$ {importacao.total_geral}"
            + (" (substituiu importação anterior dessa competência)" if existente else ""),
        )

        return jsonify({
            "message": "Folha importada com sucesso",
            "importacaoId": importacao.id,
            "qtdClt": len(colaboradores),
            "qtdPj": len(pjs),
            "totalGeral": float(importacao.total_geral),
            "substituiuAnterior": bool(existente),
        }), 201
    except Exception as e:
        db.rollback()
        print(f"[FOLHA] Erro ao importar: {e}")
        return jsonify({"message": "Erro ao salvar a importação"}), 500
    finally:
        db.close()


@bp.route("/importacoes", methods=["GET"])
@require_role_admin
def listar_importacoes():
    db = get_db()
    try:
        itens = (
            db.query(models.FolhaImportacao)
            .order_by(models.FolhaImportacao.competencia_ano.desc(), models.FolhaImportacao.competencia_mes.desc())
            .all()
        )
        return jsonify([{
            "id": i.id,
            "mes": i.competencia_mes,
            "ano": i.competencia_ano,
            "competenciaLabel": f"{MESES_PT[i.competencia_mes]}/{i.competencia_ano}",
            "arquivoNome": i.arquivo_nome,
            "importadoPor": i.importado_por,
            "totalClt": float(i.total_clt),
            "totalPj": float(i.total_pj),
            "totalGeral": float(i.total_geral),
            "qtdClt": i.qtd_clt,
            "qtdPj": i.qtd_pj,
            "createdAt": i.created_at.isoformat() if i.created_at else None,
        } for i in itens])
    finally:
        db.close()


@bp.route("/importacoes/<int:importacao_id>", methods=["DELETE"])
@require_role_admin
def apagar_importacao(importacao_id):
    db = get_db()
    try:
        importacao = db.query(models.FolhaImportacao).get(importacao_id)
        if not importacao:
            return jsonify({"message": "Importação não encontrada"}), 404
        label = f"{MESES_PT[importacao.competencia_mes]}/{importacao.competencia_ano}"
        db.delete(importacao)
        db.commit()
        audit.log(request.username, "DELETE_FOLHA", entity="folha_importacoes", entity_id=importacao_id,
                   detail=f"Apagou importação de {label}")
        return jsonify({"message": "Importação removida"})
    finally:
        db.close()


@bp.route("/resumo", methods=["GET"])
@require_role_admin
def resumo():
    db = get_db()
    try:
        importacao_id = request.args.get("importacao_id", type=int)
        if importacao_id:
            importacao = db.query(models.FolhaImportacao).get(importacao_id)
        else:
            importacao = (
                db.query(models.FolhaImportacao)
                .order_by(models.FolhaImportacao.competencia_ano.desc(), models.FolhaImportacao.competencia_mes.desc())
                .first()
            )
        if not importacao:
            return jsonify({"message": "Nenhuma folha importada ainda"}), 404

        colabs = db.query(models.FolhaColaborador).filter_by(importacao_id=importacao.id).all()
        pjs    = db.query(models.FolhaPJ).filter_by(importacao_id=importacao.id).all()

        def agrupar(chave_fn):
            grupos = {}
            for c in colabs:
                chave = chave_fn(c) or "Não informado"
                g = grupos.setdefault(chave, {"total": 0.0, "qtd": 0})
                g["total"] += float(c.total_liquido or 0)
                g["qtd"] += 1
            return [{"chave": k, "total": round(v["total"], 2), "qtd": v["qtd"]} for k, v in
                    sorted(grupos.items(), key=lambda kv: -kv[1]["total"])]

        por_regional      = agrupar(lambda c: c.regional)
        por_classificacao = agrupar(lambda c: c.classificacao)
        por_funcao        = agrupar(lambda c: c.funcao)[:10]
        por_contrato      = agrupar(lambda c: c.contrato)
        funcoes = sorted({c.funcao for c in colabs if c.funcao})

        def agrupar_pj_por_setor():
            grupos = {}
            for p in pjs:
                chave = p.setor or "Não informado"
                g = grupos.setdefault(chave, {"total": 0.0, "qtd": 0})
                g["total"] += float(p.valor or 0)
                g["qtd"] += 1
            return [{"chave": k, "total": round(v["total"], 2), "qtd": v["qtd"]} for k, v in
                    sorted(grupos.items(), key=lambda kv: -kv[1]["total"])]

        por_setor_pj = agrupar_pj_por_setor()

        rubricas = {
            "inss": sum(float(c.inss or 0) for c in colabs),
            "fgts": sum(float(c.fgts or 0) for c in colabs),
            "irpf": sum(float(c.irpf or 0) for c in colabs),
            "ajudaCusto": sum(float(c.ajuda_custo or 0) for c in colabs),
            "salarioBase": sum(float(c.salario_base or 0) for c in colabs),
        }

        # evolução: últimas até 12 competências (mais antiga → mais recente)
        historico = (
            db.query(models.FolhaImportacao)
            .order_by(models.FolhaImportacao.competencia_ano.desc(), models.FolhaImportacao.competencia_mes.desc())
            .limit(12)
            .all()
        )
        historico = list(reversed(historico))

        return jsonify({
            "importacao": {
                "id": importacao.id,
                "competenciaLabel": f"{MESES_PT[importacao.competencia_mes]}/{importacao.competencia_ano}",
                "arquivoNome": importacao.arquivo_nome,
                "importadoPor": importacao.importado_por,
                "createdAt": importacao.created_at.isoformat() if importacao.created_at else None,
            },
            "totais": {
                "geral": float(importacao.total_geral),
                "clt": float(importacao.total_clt),
                "pj": float(importacao.total_pj),
                "qtdClt": importacao.qtd_clt,
                "qtdPj": importacao.qtd_pj,
                "custoMedioClt": round(float(importacao.total_clt) / importacao.qtd_clt, 2) if importacao.qtd_clt else 0,
                "custoMedioPj": round(float(importacao.total_pj) / importacao.qtd_pj, 2) if importacao.qtd_pj else 0,
                "percentualPj": round(float(importacao.total_pj) / float(importacao.total_geral) * 100, 1) if importacao.total_geral else 0,
                "qtdColaboradores": importacao.qtd_clt + importacao.qtd_pj,
            },
            "porRegional": por_regional,
            "porClassificacao": por_classificacao,  # MOD / MOI
            "porFuncao": por_funcao,
            "porSetorPj": por_setor_pj,
            "porContrato": por_contrato,
            "funcoes": funcoes,
            "rubricas": {k: round(v, 2) for k, v in rubricas.items()},
            "pj": [{"fornecedor": p.fornecedor, "valor": float(p.valor or 0), "descricao": p.descricao, "setor": p.setor} for p in pjs],
            "evolucao": [{
                "competenciaLabel": f"{MESES_PT[h.competencia_mes]}/{h.competencia_ano}",
                "totalClt": float(h.total_clt),
                "totalPj": float(h.total_pj),
                "totalGeral": float(h.total_geral),
                "qtdClt": h.qtd_clt,
                "qtdPj": h.qtd_pj,
                "qtdColaboradores": h.qtd_clt + h.qtd_pj,
            } for h in historico],
        })
    finally:
        db.close()


@bp.route("/evolucao-headcount", methods=["GET"])
@require_role_admin
def evolucao_headcount():
    """Quantidade de colaboradores por competência (mensal), com filtro opcional de função.
    Sem filtro: CLT + PJ. Com filtro: só CLT com aquela função (PJ não tem função)."""
    db = get_db()
    try:
        funcao = (request.args.get("funcao") or "").strip() or None

        importacoes = (
            db.query(models.FolhaImportacao)
            .order_by(models.FolhaImportacao.competencia_ano.desc(), models.FolhaImportacao.competencia_mes.desc())
            .limit(24)
            .all()
        )
        importacoes = list(reversed(importacoes))

        pontos = []
        for imp in importacoes:
            if funcao:
                qtd = (
                    db.query(sa_func.count(models.FolhaColaborador.id))
                    .filter(
                        models.FolhaColaborador.importacao_id == imp.id,
                        models.FolhaColaborador.funcao == funcao,
                    )
                    .scalar()
                ) or 0
            else:
                qtd = imp.qtd_clt + imp.qtd_pj
            pontos.append({
                "competenciaLabel": f"{MESES_PT[imp.competencia_mes]}/{imp.competencia_ano}",
                "qtd": qtd,
            })

        return jsonify({"funcao": funcao, "pontos": pontos})
    finally:
        db.close()


@bp.route("/auditoria", methods=["GET"])
@require_role_admin
def auditoria():
    """Compara a competência selecionada com a competência anterior (a mais
    recente antes dela), CLT e PJ juntos: quem subiu e quem caiu de valor.
    CLT casa por nome, PJ casa por fornecedor — quem não existia na
    competência anterior (contratação nova) não aparece aqui."""
    db = get_db()
    try:
        importacao_id = request.args.get("importacao_id", type=int)
        if importacao_id:
            atual = db.query(models.FolhaImportacao).get(importacao_id)
        else:
            atual = (
                db.query(models.FolhaImportacao)
                .order_by(models.FolhaImportacao.competencia_ano.desc(), models.FolhaImportacao.competencia_mes.desc())
                .first()
            )
        if not atual:
            return jsonify({"message": "Nenhuma folha importada ainda"}), 404

        anterior = (
            db.query(models.FolhaImportacao)
            .filter(sa_or(
                models.FolhaImportacao.competencia_ano < atual.competencia_ano,
                sa_and(
                    models.FolhaImportacao.competencia_ano == atual.competencia_ano,
                    models.FolhaImportacao.competencia_mes < atual.competencia_mes,
                ),
            ))
            .order_by(models.FolhaImportacao.competencia_ano.desc(), models.FolhaImportacao.competencia_mes.desc())
            .first()
        )

        competencia_atual_label = f"{MESES_PT[atual.competencia_mes]}/{atual.competencia_ano}"

        if not anterior:
            return jsonify({
                "competenciaAtual": competencia_atual_label,
                "competenciaAnterior": None,
                "positivas": [],
                "negativas": [],
                "message": "Não há competência anterior para comparar",
            })

        colabs_atual = db.query(models.FolhaColaborador).filter_by(importacao_id=atual.id).all()
        colabs_ant   = db.query(models.FolhaColaborador).filter_by(importacao_id=anterior.id).all()
        pjs_atual    = db.query(models.FolhaPJ).filter_by(importacao_id=atual.id).all()
        pjs_ant      = db.query(models.FolhaPJ).filter_by(importacao_id=anterior.id).all()

        mapa_ant_clt = {_normaliza_nome(c.nome): c for c in colabs_ant}
        mapa_ant_pj  = {_normaliza_nome(p.fornecedor): p for p in pjs_ant}

        variacoes = []
        for c in colabs_atual:
            ant = mapa_ant_clt.get(_normaliza_nome(c.nome))
            if not ant:
                continue
            delta = round(float(c.total_liquido or 0) - float(ant.total_liquido or 0), 2)
            if delta == 0:
                continue
            variacoes.append({
                "tipo": "CLT",
                "nome": c.nome,
                "referencia": c.funcao,
                "valorAnterior": float(ant.total_liquido or 0),
                "valorAtual": float(c.total_liquido or 0),
                "delta": delta,
                "percentual": round(delta / float(ant.total_liquido) * 100, 1) if ant.total_liquido else None,
            })

        for p in pjs_atual:
            ant = mapa_ant_pj.get(_normaliza_nome(p.fornecedor))
            if not ant:
                continue
            delta = round(float(p.valor or 0) - float(ant.valor or 0), 2)
            if delta == 0:
                continue
            variacoes.append({
                "tipo": "PJ",
                "nome": p.fornecedor,
                "referencia": p.setor,
                "valorAnterior": float(ant.valor or 0),
                "valorAtual": float(p.valor or 0),
                "delta": delta,
                "percentual": round(delta / float(ant.valor) * 100, 1) if ant.valor else None,
            })

        positivas = sorted([v for v in variacoes if v["delta"] > 0], key=lambda v: -v["delta"])
        negativas = sorted([v for v in variacoes if v["delta"] < 0], key=lambda v: v["delta"])

        return jsonify({
            "competenciaAtual": competencia_atual_label,
            "competenciaAnterior": f"{MESES_PT[anterior.competencia_mes]}/{anterior.competencia_ano}",
            "positivas": positivas,
            "negativas": negativas,
        })
    finally:
        db.close()
